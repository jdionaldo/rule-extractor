#!/usr/bin/env python3
"""Generate a small sample policy .docx so you can try the pipeline immediately.

    python make_sample.py            # writes ./sample_docs/access_policy.docx

The sample deliberately contains a couple of conflicting rules so the
conflict-detection stage has something to find.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

PARAGRAPHS = [
    ("Heading 1", "Access Control Policy"),
    ("Heading 2", "1. Authentication"),
    ("Normal", "All employees must enable multi-factor authentication on their "
               "corporate accounts before accessing internal systems."),
    ("Normal", "Passwords must be at least 12 characters long and rotated every "
               "90 days."),
    ("Normal", "Contractors may access the guest network without multi-factor "
               "authentication for the duration of their engagement."),
    ("Heading 2", "2. Data Handling"),
    ("Normal", "Customer data must be encrypted at rest using AES-256."),
    ("Normal", "Employees should not store customer data on personal devices."),
    ("Normal", "Backups of customer data must be retained for at least 7 years."),
    ("Heading 2", "3. Retention"),
    ("Normal", "Customer data must be deleted within 30 days of account closure."),
    ("Normal", "All access logs must be reviewed weekly by the security team."),
]


def main() -> None:
    out_dir = Path(__file__).resolve().parent / "sample_docs"
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for style, text in PARAGRAPHS:
        try:
            doc.add_paragraph(text, style=style)
        except KeyError:
            doc.add_paragraph(text)
    dest = out_dir / "access_policy.docx"
    doc.save(str(dest))
    print(f"Wrote {dest}")


if __name__ == "__main__":
    main()
