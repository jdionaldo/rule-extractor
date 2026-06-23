"""Flexible document ingestion.

Reads many document formats into a single, uniform list of verbatim ``Span``s
so the rest of the pipeline (extraction, conflict detection, the UI) is
format-agnostic. To support a new format, add a reader function and register it
in ``_READERS`` — nothing downstream needs to change.

Supported today: .docx, .pdf, .pptx, .txt, .md

Every reader emits ``(text, location, heading)`` tuples; ``_build_spans`` turns
those into ``Span`` objects with a stable, sequential id per document.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Callable, List, Optional, Tuple

from .models import Span

# A unit of extracted text plus where it came from and its section heading.
Item = Tuple[str, Optional[str], Optional[str]]
Reader = Callable[[Path], List["Item"]]


class IngestError(RuntimeError):
    """Raised when a file cannot be read (corrupt, encrypted, unsupported)."""


# --- shared helpers ---------------------------------------------------------


def _split_blocks(text: str) -> List[str]:
    """Split free text into paragraph-like blocks on blank lines.

    Single newlines inside a block are collapsed to spaces so wrapped lines
    rejoin into readable sentences.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks: List[str] = []
    for raw in re.split(r"\n\s*\n", text):
        joined = " ".join(line.strip() for line in raw.split("\n") if line.strip())
        joined = joined.strip()
        if joined:
            blocks.append(joined)
    return blocks


def _build_spans(doc: str, items: List[Item]) -> List[Span]:
    spans: List[Span] = []
    for i, (text, location, heading) in enumerate(items):
        text = (text or "").strip()
        if not text:
            continue
        spans.append(
            Span(
                span_id=f"{doc}::{i:04d}",
                doc=doc,
                index=i,
                text=text,
                heading=heading,
                location=location,
            )
        )
    return spans


def _is_heading_style(style_name: str | None) -> bool:
    return bool(style_name) and style_name.lower().startswith(("heading", "title"))


# --- per-format readers -----------------------------------------------------


def _read_docx(path: Path) -> List[Item]:
    from docx import Document

    doc = Document(str(path))
    items: List[Item] = []
    current_heading: str | None = None
    for index, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if not text:
            continue
        style_name = getattr(getattr(para, "style", None), "name", None)
        if _is_heading_style(style_name):
            current_heading = text
        items.append((text, f"paragraph {index}", current_heading))
    return items


def _read_pdf(path: Path) -> List[Item]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            # Try an empty-password unlock; many PDFs are "encrypted" with none.
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                raise IngestError(f"{path.name} is password-protected.")
        pages = reader.pages
    except IngestError:
        raise
    except Exception as exc:  # noqa: BLE001 - surface a clean message
        raise IngestError(f"Could not read PDF {path.name}: {exc}") from exc

    items: List[Item] = []
    for page_no, page in enumerate(pages, start=1):
        text = page.extract_text() or ""
        for block in _split_blocks(text):
            items.append((block, f"page {page_no}", None))
    return items


def _read_pptx(path: Path) -> List[Item]:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"Could not read PPTX {path.name}: {exc}") from exc

    items: List[Item] = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        title = None
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except Exception:  # noqa: BLE001 - some layouts have no title placeholder
            title = None
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    items.append((text, f"slide {slide_no}", title))
    return items


def _read_text(path: Path) -> List[Item]:
    content = path.read_text(encoding="utf-8", errors="replace")
    return [
        (block, f"block {i + 1}", None)
        for i, block in enumerate(_split_blocks(content))
    ]


# Extension -> reader. The single source of truth for "what can we ingest".
_READERS: dict[str, Reader] = {
    ".docx": _read_docx,
    ".pdf": _read_pdf,
    ".pptx": _read_pptx,
    ".txt": _read_text,
    ".md": _read_text,
}

SUPPORTED_EXTENSIONS = tuple(sorted(_READERS))


# --- public API -------------------------------------------------------------


def read_file(path: Path) -> List[Span]:
    """Read a single supported document into Spans."""
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        raise IngestError(
            f"Unsupported file type: {path.name} "
            f"(supported: {', '.join(SUPPORTED_EXTENSIONS)})"
        )
    return _build_spans(path.name, reader(path))


def collect_paths(input_path: Path, limit: int | None = None) -> List[Path]:
    """Resolve a file or directory to a sorted list of supported documents."""
    if input_path.is_file():
        paths = (
            [input_path] if input_path.suffix.lower() in _READERS else []
        )
    else:
        paths = sorted(
            p
            for p in input_path.rglob("*")
            if p.is_file()
            and p.suffix.lower() in _READERS
            and not p.name.startswith("~$")  # skip Office lock files
        )
    if limit is not None:
        paths = paths[:limit]
    return paths


def read_all(input_path: Path, limit: int | None = None) -> List[Span]:
    """Read every supported document under ``input_path`` into a flat list."""
    spans: List[Span] = []
    for path in collect_paths(input_path, limit=limit):
        spans.extend(read_file(path))
    return spans
